# -*- coding: utf-8 -*-
"""Dựng file Word "Report 6 – Hướng dẫn sử dụng phần mềm" (mẫu SEP490) từ hướng dẫn trong app.

Phần Hướng dẫn sử dụng (mục II.3.2 trở đi) KHÔNG viết lại ở đây: script đọc
thẳng 6 trang web/jsp/guide/*.jsp -- đúng các trang /guide người dùng mở trong
app -- theo thứ tự tab ở _top.jspf, rồi chuyển từng thẻ sang Word. Sửa hướng
dẫn trong app xong thì chạy lại là bản Word theo kịp, không có bản thứ hai để
lệch nhau.

Phần còn lại (lịch sử thay đổi, gói bàn giao, cài đặt, tổng quan) viết ở cuối
file này, theo README.md, DEPLOY.md, db/ và PERMISSIONS.md.

Chạy:
    python tools/guide/word/export_word.py --template "<Report6_Software User Guides.docx>" --out "<file.docx>"

--template là file mẫu Report6 của môn học: lấy bìa, style, khổ giấy; file mẫu
không bị sửa. Cần python-docx, beautifulsoup4, Pillow. Ba sơ đồ *.png cạnh
file này do diagrams.py vẽ.

Trang hướng dẫn chỉ dùng tập thẻ ghi ở đầu web/css/guide.css. Thêm thẻ hay lớp
mới ở đó thì dạy thêm cho block() / inline_runs() bên dưới: gặp thẻ lạ, script
dừng và báo tên thẻ, chứ không lặng lẽ bỏ mất nội dung.
"""
import argparse
import datetime
import re
from pathlib import Path

from bs4 import BeautifulSoup, Comment, NavigableString, Tag
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Emu, Pt, RGBColor
from docx.text.paragraph import Paragraph
from PIL import Image

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[2]
JSP_DIR = ROOT / 'web' / 'jsp' / 'guide'
IMG_DIR = ROOT / 'web' / 'guide'

PROJECT_NAME = ('POSCS - Hệ thống Quản lý Dữ liệu Khách hàng và Dịch vụ Hỗ trợ Kỹ thuật POSTEF '
                '– Chi nhánh Miền Bắc')
REPO_URL = 'https://github.com/Bunky2k3/POSCS.git'
CHAPTER = 2   # "Hình 2-N": mọi hình nằm trong chương II, đánh số như Project_Report ("Hình 4-27")

MARK = RGBColor(0xE1, 0x1D, 0x48)        # --g-mark của guide.css: số khoanh đỏ trên ảnh
MARK_PLAIN = RGBColor(0x9C, 0xA3, 0xAF)  # li.plain: bước không khoanh trên ảnh
NAVY = RGBColor(0x00, 0x3C, 0x6E)
MUTED = RGBColor(0x6B, 0x72, 0x80)
HEAD_FILL = 'FFE8E1'                     # nền hàng tiêu đề bảng, theo bảng của file mẫu
CALLOUT = {'guide-note': ('EEF6FD', '0F9EDB', 'Lưu ý'), 'guide-warn': ('FFF6E6', 'F5A623', 'Chú ý')}
# Biểu tượng Font Awesome trong câu chữ ("bấm <i class="fa-solid fa-pen">") -> ký tự gần nghĩa.
ICONS = {'fa-eye': '👁', 'fa-pen': '✎', 'fa-trash': '🗑', 'fa-lock': '🔒'}
SYMBOL_FONT = 'Segoe UI Symbol'
BLOCK_TAGS = {'p', 'ul', 'ol', 'div', 'table', 'figure', 'h1', 'h2', 'h3'}

# Sơ đồ chèn vào giữa trang hướng dẫn: module -> (ảnh, chú thích, chèn NGAY TRƯỚC tiêu đề có id này).
DIAGRAMS = {
    'contract': ('vong-doi-hop-dong.png', 'Vòng đời hợp đồng: tiến độ và trạng thái theo lịch', 'quyen'),
    'ticket': ('trang-thai-phieu.png', 'Trạng thái và bốn phần nội dung của phiếu hỗ trợ', 'quyen'),
}
OVERVIEW = ('tong-quan.png', 'Tổng quan các phân hệ, vai trò quản lý và luồng dữ liệu')
# Câu chỉ đúng trên trang web (nói tới thanh tab của trang Hướng dẫn) -> câu cho bản giấy.
WEB_ONLY = {
    'general': [('Hướng dẫn từng phân hệ nằm ở các tab bên cạnh.',
                 'Hướng dẫn từng phân hệ nằm ở các mục tiếp theo.')],
}


def fail(msg):
    raise SystemExit('export_word: ' + msg)


# ----------------------------------------------------------------------
# Đọc trang hướng dẫn
# ----------------------------------------------------------------------

def tabs():
    """[(module, tên tab)] đúng thứ tự tab trong app."""
    top = (JSP_DIR / '_top.jspf').read_text(encoding='utf-8')
    found = re.findall(r'guide\?module=(\w+)" class="[^"]*">([^<]+)</a>', top)
    if not found:
        fail('không đọc được thanh tab trong _top.jspf')
    return found


def load(module):
    src = (JSP_DIR / (module + '.jsp')).read_text(encoding='utf-8')
    src = re.sub(r'<%--.*?--%>', '', src, flags=re.S)
    src = re.sub(r'<%@.*?%>', '', src, flags=re.S)
    src = re.sub(r'<c:set\b[^>]*/>', '', src)
    if '<%' in src or re.search(r'<c:|<fmt:|\$\{(?!img\})', src):
        fail(module + '.jsp có mã JSP ngoài ${img}: dạy thêm cho load()')
    for web, word in WEB_ONLY.get(module, []):
        if web not in src:
            fail('%s.jsp đã đổi câu "%s": sửa lại WEB_ONLY' % (module, web))
        src = src.replace(web, word)
    return BeautifulSoup(src.replace('${img}', 'IMG:' + module), 'html.parser')


class Module:
    """Một trang hướng dẫn = một mục 3.x. Đánh số mục và số hình trước khi dựng."""

    def __init__(self, key, tab, num):
        self.key, self.tab, self.num = key, tab, num
        self.soup = load(key)
        h1 = self.soup.find('h1')
        self.title = h1.get_text(' ', strip=True) if h1 else tab
        self.sections = {}    # id tiêu đề -> số mục trong Word ("3.2.1")
        self.by_number = {}   # số mục trong trang ("5.1") -> số mục trong Word
        h2 = h3 = 0
        for h in self.soup.find_all(['h2', 'h3']):
            if h.name == 'h2':
                h2, h3 = h2 + 1, 0
                own = str(h2)
            else:
                h3 += 1
                own = '%d.%d' % (h2, h3)
            m = re.match(r'([\d.]+?)\.?\s', h.get_text())
            if not m or m.group(1) != own:
                fail('%s.jsp: tiêu đề "%s" lẽ ra mang số %s' % (key, h.get_text(strip=True), own))
            if not h.get('id'):
                fail('%s.jsp: tiêu đề "%s" thiếu id' % (key, h.get_text(strip=True)))
            self.sections[h['id']] = self.num + '.' + own
            self.by_number[own] = self.num + '.' + own
        self.figs = {}        # số hình trong trang -> số hình trong Word
        self.diagram_fig = None

    def number_figures(self, counter):
        """Cấp số hình theo đúng thứ tự xuất hiện, kể cả sơ đồ chèn thêm. Trả về bộ đếm mới."""
        diagram = DIAGRAMS.get(self.key)
        local = 0
        for node in self.soup.find_all(['figure', 'h2', 'h3']):
            if diagram and node.get('id') == diagram[2]:
                counter += 1
                self.diagram_fig = counter
            if node.name == 'figure':
                local += 1
                cap = node.find('figcaption')
                m = re.match(r'Hình (\d+)\.', cap.get_text(strip=True) if cap else '')
                if not m or int(m.group(1)) != local:
                    fail('%s.jsp: hình thứ %d có chú thích "%s"' % (self.key, local, cap and cap.get_text()))
                counter += 1
                self.figs[local] = counter
        if diagram and self.diagram_fig is None:
            fail('%s.jsp: không có tiêu đề id="%s" để chèn sơ đồ' % (self.key, diagram[2]))
        return counter


class Book:
    def __init__(self):
        self.modules = [Module(k, t, '3.%d' % (i + 2)) for i, (k, t) in enumerate(tabs())]
        counter = 1   # Hình 2-1 là sơ đồ tổng quan ở mục 3.1
        for m in self.modules:
            counter = m.number_figures(counter)
        self.fig_total = counter

    def module(self, key):
        return next(m for m in self.modules if m.key == key)

    def section(self, key, anchor):
        return self.module(key).sections[anchor]

    def rewrite(self, mod, text):
        """Số hình, số mục viết tay trong câu chữ -> số trong Word."""
        def fig(m):
            n = int(m.group(1))
            if n not in mod.figs:
                fail('%s.jsp nhắc tới Hình %d nhưng trang không có hình đó' % (mod.key, n))
            return 'Hình %d-%d' % (CHAPTER, mod.figs[n])

        def other(m):
            target = next((x for x in self.modules if x.tab == m.group(1).strip()), None)
            if target is None or m.group(2) not in target.by_number:
                fail('%s.jsp: không tìm thấy "hướng dẫn %s, mục %s"' % (mod.key, m.group(1), m.group(2)))
            return 'mục ' + target.by_number[m.group(2)]

        text = re.sub(r'Hình (\d+)', fig, text)
        return re.sub(r'hướng dẫn ([^,()]+), mục (\d+(?:\.\d+)*)', other, text)

    def link_text(self, mod, a):
        """Link trong trang không bấm được trên giấy: đổi thành số mục của bản Word."""
        text = re.sub(r'\s+', ' ', a.get_text())
        href = a.get('href', '')
        if not href.startswith('#'):
            fail('%s.jsp: link ra ngoài trang (%s): dạy thêm cho link_text()' % (mod.key, href))
        if href[1:] not in mod.sections:
            fail('%s.jsp: link tới #%s nhưng không có tiêu đề đó' % (mod.key, href[1:]))
        num = mod.sections[href[1:]]
        if re.fullmatch(r'mục [\d.]+', text.strip()):       # "mục 5"
            return 'mục ' + num
        if re.fullmatch(r'[\d.]+', text.strip()):           # "mục <a>6</a>, <a>12</a>"
            return num
        if re.fullmatch(r'Hình \d+', text.strip()):         # "giống <a>Hình 2</a>"
            return self.rewrite(mod, text)
        return '%s (mục %s)' % (self.rewrite(mod, text), num)   # "lập <a>phụ lục</a>"


# ----------------------------------------------------------------------
# Tiện ích OOXML
# ----------------------------------------------------------------------

def set_font(run, name):
    """Đặt font cho mọi loại ký tự: ký hiệu (❶, ✎) Word có thể xếp vào nhóm Đông Á."""
    run.font.name = name
    rfonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    rfonts.set(qn('w:eastAsia'), name)
    rfonts.set(qn('w:cs'), name)


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    # Thứ tự con của tcPr do schema quy định: tcW, gridSpan, vMerge, tcBorders, shd, ...
    after = tcPr.find(qn('w:tcBorders'))
    if after is None:
        after = tcPr.find(qn('w:tcW'))
    if after is not None:
        after.addnext(shd)
    else:
        tcPr.insert(0, shd)


def cell_borders(cell, **sides):
    """sides: left=('single', 24, 'F5A623'), top=None (không viền)..."""
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement('w:tcBorders')
    for side in ('top', 'left', 'bottom', 'right'):
        el = OxmlElement('w:' + side)
        spec = sides.get(side)
        if spec:
            el.set(qn('w:val'), spec[0])
            el.set(qn('w:sz'), str(spec[1]))
            el.set(qn('w:space'), '0')
            el.set(qn('w:color'), spec[2])
        else:
            el.set(qn('w:val'), 'nil')
        borders.append(el)
    tcW = tcPr.find(qn('w:tcW'))
    if tcW is not None:
        tcW.addnext(borders)
    else:
        tcPr.insert(0, borders)


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement('w:tblHeader')
    el.set(qn('w:val'), 'true')
    trPr.append(el)


def keep_row_whole(row):
    """Không cắt dòng bảng qua hai trang (khối lệnh, hộp lưu ý). Dài hơn một trang thì Word vẫn cắt."""
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement('w:cantSplit')
    trPr.insert(0, el)   # schema: cantSplit đứng trước trHeight, tblHeader


def page_break_before(paragraph):
    paragraph.paragraph_format.page_break_before = True


def step_marker(n):
    """Số tròn nền đặc giống số khoanh trên ảnh: ❶..❿ rồi ⓫..⓴."""
    if 1 <= n <= 10:
        return chr(0x2776 + n - 1)
    if 11 <= n <= 20:
        return chr(0x24EB + n - 11)
    return '(%d)' % n


# ----------------------------------------------------------------------
# Dựng tài liệu
# ----------------------------------------------------------------------

class Writer:
    def __init__(self, template, book):
        self.doc = Document(str(template))
        self.book = book
        sec = self.doc.sections[-1]
        self.text_w = Emu(int(sec.page_width - sec.left_margin - sec.right_margin))
        self.text_h = Emu(int(sec.page_height - sec.top_margin - sec.bottom_margin))
        self.bullet_num = None

    # --- khối cơ bản ---------------------------------------------------

    def heading(self, text, level, new_page=False):
        p = self.doc.add_paragraph(text, style='Heading %d' % level)
        if new_page:
            page_break_before(p)
        return p

    def spacer(self, box):
        """Đoạn trống thấp sau bảng: Word dán đoạn kế tiếp sát mép bảng."""
        p = box.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = Pt(6)

    def emit(self, p, runs):
        """Ghi các (chữ, định dạng) vào đoạn p, gộp khoảng trắng như trình duyệt."""
        clean = []
        for text, fmt in runs:
            if clean and text.startswith(' ') and clean[-1][0].endswith((' ', '\n')):
                text = text[1:]
            if text:
                clean.append([text, fmt])
        if clean:
            clean[0][0] = clean[0][0].lstrip(' ')
            clean[-1][0] = clean[-1][0].rstrip(' ')
        for text, fmt in clean:
            if not text:
                continue
            for i, part in enumerate(text.split('\n')):
                if i:
                    p.add_run().add_break()
                if not part:
                    continue
                r = p.add_run(part)
                r.bold = True if 'b' in fmt else None
                r.italic = True if 'i' in fmt else None
                r.underline = True if 'u' in fmt else None
                if 'code' in fmt:
                    set_font(r, 'Consolas')
                    r.font.size = Pt(9.5)
                if 'icon' in fmt:
                    set_font(r, SYMBOL_FONT)

    def inline_runs(self, nodes, mod, fmt=frozenset()):
        out = []
        for ch in nodes:
            if isinstance(ch, Comment):
                continue
            if isinstance(ch, NavigableString):
                text = re.sub(r'\s+', ' ', str(ch))
                out.append((self.book.rewrite(mod, text) if mod else text, fmt))
            elif ch.name in ('strong', 'b'):
                out += self.inline_runs(ch.children, mod, fmt | {'b'})
            elif ch.name == 'em':
                out += self.inline_runs(ch.children, mod, fmt | {'i'})
            elif ch.name == 'u':
                out += self.inline_runs(ch.children, mod, fmt | {'u'})
            elif ch.name == 'code':
                out += self.inline_runs(ch.children, mod, fmt | {'code'})
            elif ch.name == 'i':
                icon = next((ICONS[c] for c in ch.get('class', []) if c in ICONS), None)
                if icon is None:
                    fail('biểu tượng chưa có trong ICONS: %s' % ch.get('class'))
                out.append((icon, fmt | {'icon'}))
            elif ch.name == 'a':
                out.append((self.book.link_text(mod, ch), fmt))
            elif ch.name == 'span':
                out += self.inline_runs(ch.children, mod, fmt)
            elif ch.name == 'br':
                out.append(('\n', fmt))
            else:
                fail('%s.jsp: thẻ <%s> trong dòng chữ: dạy thêm cho inline_runs()' % (mod.key, ch.name))
        return out

    # --- các khối của trang hướng dẫn ------------------------------------

    def block(self, node, mod, box, width):
        if isinstance(node, Comment):
            return
        if isinstance(node, NavigableString):
            if str(node).strip():
                fail('%s.jsp: chữ nằm ngoài thẻ: "%s"' % (mod.key, str(node).strip()[:60]))
            return
        cls = node.get('class', [])
        if node.name == 'h1':
            return   # tên trang đã thành tiêu đề mục 3.x
        if node.name in ('h2', 'h3'):
            title = re.sub(r'^[\d.]+\s*', '', node.get_text(' ', strip=True))
            self.heading(mod.sections[node['id']] + ' ' + title, 4 if node.name == 'h2' else 5)
        elif node.name == 'p':
            p = box.add_paragraph()
            runs = self.inline_runs(node.children, mod)
            if 'lead' in cls:
                runs = [(t, f | {'i'}) for t, f in runs]
            self.emit(p, runs)
            nxt = node.find_next_sibling()
            if nxt is not None and nxt.name == 'figure' and len(node.get_text()) < 300:
                p.paragraph_format.keep_with_next = True   # câu dẫn ngắn đi cùng trang với hình
        elif node.name == 'div' and 'guide-who' in cls:
            self.who(node, box)
        elif node.name == 'div' and set(cls) & set(CALLOUT):
            self.callout(node, mod, box, width, CALLOUT[(set(cls) & set(CALLOUT)).pop()])
        elif node.name == 'figure' and 'guide-shot' in cls:
            img = node.find('img')
            src = img['src']
            if not src.startswith('IMG:'):
                fail('%s.jsp: ảnh không nằm dưới ${img}: %s' % (mod.key, src))
            module, name = src[4:].split('/', 1)
            cap = re.sub(r'^Hình \d+\.\s*', '', node.find('figcaption').get_text(' ', strip=True))
            local = len(node.find_all_previous('figure')) + 1
            self.picture(box, IMG_DIR / module / name, mod.figs[local], cap, width,
                         narrow='narrow' in cls, border=True)
        elif node.name == 'ol' and 'guide-steps' in cls:
            self.steps(node, mod, box, width)
        elif node.name == 'ul':
            self.bullets(node, mod, box, width)
        elif node.name == 'table' and 'guide-table' in cls:
            rows = []
            for tr in node.find_all('tr'):
                cells = tr.find_all(['th', 'td'], recursive=False)
                if any(c.get('colspan') or c.get('rowspan') for c in cells):
                    fail('%s.jsp: bảng có colspan/rowspan: dạy thêm cho table()' % mod.key)
                rows.append([(c.name == 'th', self.inline_runs(c.children, mod)) for c in cells])
            self.table(box, rows, width)
        else:
            fail('%s.jsp: khối <%s class="%s">: dạy thêm cho block()' % (mod.key, node.name, ' '.join(cls)))

    def who(self, node, box):
        p = box.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True   # dòng này luôn nằm giữa tiêu đề và nội dung
        lbl = node.find('span', class_='lbl')
        r = p.add_run((lbl.get_text(strip=True) if lbl else 'Ai làm được:') + ' ')
        r.bold, r.font.size, r.font.color.rgb = True, Pt(10), MUTED
        for i, span in enumerate(node.find_all('span', class_='role')):
            if i:
                sep = p.add_run('  ·  ')
                sep.font.size, sep.font.color.rgb = Pt(10), MUTED
            r = p.add_run(span.get_text(strip=True))
            r.bold, r.font.size = True, Pt(10)
            if 'no' in span.get('class', []):
                r.font.strike, r.font.color.rgb = True, MARK_PLAIN
            else:
                r.font.color.rgb = NAVY

    def callout(self, node, mod, box, width, style):
        fill, edge, label = style
        t = box.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = t.cell(0, 0)
        cell.width = Emu(int(width))
        cell_borders(cell, left=('single', 24, edge))
        shade(cell, fill)
        keep_row_whole(t.rows[0])
        first = True
        for ch in node.children:
            if isinstance(ch, NavigableString) and not str(ch).strip():
                continue
            if isinstance(ch, Tag) and ch.name == 'p':
                p = cell.paragraphs[0] if first else cell.add_paragraph()
                p.paragraph_format.space_after = Pt(4)
                runs = self.inline_runs(ch.children, mod)
                if first:
                    runs = [(label + ': ', frozenset({'b', 'label'}))] + runs
                self.emit(p, runs)
                if first:
                    p.runs[0].font.color.rgb = RGBColor.from_string(edge)
            else:
                if first:
                    lp = cell.paragraphs[0]
                    r = lp.add_run(label + ':')
                    r.bold, r.font.color.rgb = True, RGBColor.from_string(edge)
                self.block(ch, mod, cell, width - Cm(0.6))
            first = False
        self.spacer(box)

    def picture(self, box, path, number, caption, width, narrow=False, border=False):
        w_px, h_px = Image.open(path).size
        w = width * (0.72 if narrow else 1.0)
        w = min(w, Emu(int(w_px / 110 * 914400)))    # không phóng ảnh nhỏ quá ~110 dpi
        if w * h_px / w_px > self.text_h * 0.78:        # ảnh dọc cao quá: bó theo chiều cao
            w = self.text_h * 0.78 * w_px / h_px
        p = box.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_after = Pt(4)
        shape = p.add_run().add_picture(str(path), width=Emu(int(w)))
        if border:   # ảnh chụp nền trắng: viền mảnh cho khỏi lẫn vào trang giấy
            shape._inline.graphic.graphicData.pic.spPr.append(parse_xml(
                '<a:ln %s w="9525"><a:solidFill><a:srgbClr val="D1D5DB"/></a:solidFill></a:ln>' % nsdecls('a')))
        c = box.add_paragraph()
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        c.paragraph_format.space_after = Pt(12)
        r = c.add_run('Hình %d-%d. %s' % (CHAPTER, number, caption))
        r.italic, r.font.size, r.font.color.rgb = True, Pt(10), MUTED

    def steps(self, ol, mod, box, width):
        for n, li in enumerate(ol.find_all('li', recursive=False), 1):
            p = box.add_paragraph()
            pf = p.paragraph_format
            pf.left_indent, pf.first_line_indent, pf.space_after = Cm(0.9), Cm(-0.9), Pt(4)
            pf.tab_stops.add_tab_stop(Cm(0.9))
            r = p.add_run(step_marker(n))
            set_font(r, SYMBOL_FONT)
            r.font.size = Pt(12)
            r.font.color.rgb = MARK_PLAIN if 'plain' in li.get('class', []) else MARK
            p.add_run('\t')
            inline = [c for c in li.children if not (isinstance(c, Tag) and c.name in BLOCK_TAGS)]
            self.emit(p, self.inline_runs(inline, mod))
            for c in li.children:
                if isinstance(c, Tag) and c.name == 'ul':
                    self.bullets(c, mod, box, width, indent=Cm(0.9))
                elif isinstance(c, Tag) and c.name in BLOCK_TAGS:
                    fail('%s.jsp: khối <%s> trong một bước: dạy thêm cho steps()' % (mod.key, c.name))

    def bullets(self, ul, mod, box, width, indent=Cm(0)):
        num = self.bullet_numbering()
        for li in ul.find_all('li', recursive=False):
            p = box.add_paragraph(style='List Paragraph')
            numPr = p._p.get_or_add_pPr().get_or_add_numPr()
            numPr.get_or_add_ilvl().val = 0
            numPr.get_or_add_numId().val = num
            p.paragraph_format.left_indent = Emu(int(indent + Cm(0.63)))
            p.paragraph_format.first_line_indent = Cm(-0.63)
            inline = [c for c in li.children if not (isinstance(c, Tag) and c.name in BLOCK_TAGS)]
            self.emit(p, self.inline_runs(inline, mod) if mod else [(li.get_text(), frozenset())])
            for c in li.children:
                if isinstance(c, Tag) and c.name in BLOCK_TAGS:
                    fail('%s.jsp: khối <%s> trong một dòng gạch đầu dòng' % (mod.key, c.name))

    def bullet_numbering(self):
        """Một định nghĩa gạch đầu dòng dùng chung, thêm vào numbering.xml của file mẫu."""
        if self.bullet_num is None:
            numbering = self.doc.part.numbering_part.element
            abs_ids = [int(a.get(qn('w:abstractNumId'))) for a in numbering.findall(qn('w:abstractNum'))]
            num_ids = [int(n.get(qn('w:numId'))) for n in numbering.findall(qn('w:num'))]
            abs_id, num_id = max(abs_ids + [0]) + 1, max(num_ids + [0]) + 1
            abstract = parse_xml(
                '<w:abstractNum %s w:abstractNumId="%d"><w:multiLevelType w:val="singleLevel"/>'
                '<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="bullet"/><w:lvlText w:val="•"/>'
                '<w:lvlJc w:val="left"/><w:pPr><w:ind w:left="357" w:hanging="357"/></w:pPr>'
                '<w:rPr><w:rFonts w:ascii="Calibri" w:hAnsi="Calibri"/></w:rPr></w:lvl></w:abstractNum>'
                % (nsdecls('w'), abs_id))
            # Schema: mọi abstractNum đứng trước mọi num.
            first_num = numbering.find(qn('w:num'))
            if first_num is not None:
                first_num.addprevious(abstract)
            else:
                numbering.append(abstract)
            numbering.append(parse_xml('<w:num %s w:numId="%d"><w:abstractNumId w:val="%d"/></w:num>'
                                       % (nsdecls('w'), num_id, abs_id)))
            self.bullet_num = num_id
        return self.bullet_num

    def table(self, box, rows, width, widths=None):
        """rows: [[(là_tiêu_đề, runs)]]. Bảng theo kiểu bảng của file mẫu."""
        ncol = max(len(r) for r in rows)
        if widths is None:
            lens = []
            for j in range(ncol):
                cells = [sum(len(t) for t, _ in r[j][1]) for r in rows if j < len(r)]
                lens.append(min(80, max(10, sum(cells) / len(cells))))
            widths = [width * x / sum(lens) for x in lens]
        t = box.add_table(rows=len(rows), cols=ncol)
        t.style = self.doc.styles['Table Grid']
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        for j, w in enumerate(widths):
            t.columns[j].width = Emu(int(w))
        for i, row in enumerate(rows):
            head_row = all(h for h, _ in row)
            for j in range(ncol):
                cell = t.cell(i, j)
                cell.width = Emu(int(widths[j]))
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                p = cell.paragraphs[0]
                if j < len(row):
                    head, runs = row[j]
                    p.style = self.doc.styles['Heading Lv1' if head else 'Bang']
                    self.emit(p, runs)
                    if head:
                        shade(cell, HEAD_FILL)
            if head_row:
                repeat_header(t.rows[i])
        self.spacer(box)
        return t

    def plain_table(self, header, body, widths_cm):
        rows = [[(True, [(h, frozenset())]) for h in header]]
        for r in body:
            rows.append([(False, c if isinstance(c, list) else [(c, frozenset())]) for c in r])
        total = sum(widths_cm)
        return self.table(self.doc, rows, self.text_w, [self.text_w * w / total for w in widths_cm])

    def code(self, lines, size=9):
        # Câu dẫn ngay trên đi cùng trang với khối lệnh, và khối lệnh không bị cắt ngang trang.
        self.doc.paragraphs[-1].paragraph_format.keep_with_next = True
        t = self.doc.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = t.cell(0, 0)
        cell.width = self.text_w
        cell_borders(cell, left=('single', 18, '9CA3AF'))
        shade(cell, 'F3F4F6')
        keep_row_whole(t.rows[0])
        for i, line in enumerate(lines):
            p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            r = p.add_run(line)
            set_font(r, 'Consolas')
            r.font.size = Pt(size)
        self.spacer(self.doc)

    def rich(self, text):
        """Câu chữ viết tay trong file này: **đậm**, `mã`. Trả về runs cho emit()."""
        out = []
        for part in re.split(r'(\*\*[^*]+\*\*|`[^`]+`)', text):
            if part.startswith('**'):
                out.append((part[2:-2], frozenset({'b'})))
            elif part.startswith('`'):
                out.append((part[1:-1], frozenset({'code'})))
            elif part:
                out.append((part, frozenset()))
        return out

    def p(self, text, box=None):
        p = (box or self.doc).add_paragraph()
        self.emit(p, self.rich(text))
        return p

    def step_title(self, text):
        """Tiêu đề nhỏ in đậm ("Bước 1. ..."), luôn đi cùng trang với đoạn ngay dưới."""
        p = self.p('**%s**' % text)
        p.paragraph_format.keep_with_next = True
        return p

    def bullet_list(self, items):
        num = self.bullet_numbering()
        for text in items:
            p = self.doc.add_paragraph(style='List Paragraph')
            numPr = p._p.get_or_add_pPr().get_or_add_numPr()
            numPr.get_or_add_ilvl().val = 0
            numPr.get_or_add_numId().val = num
            p.paragraph_format.left_indent = Cm(0.63)
            p.paragraph_format.first_line_indent = Cm(-0.63)
            self.emit(p, self.rich(text))

    # --- khung tài liệu -------------------------------------------------

    def prepare(self, today):
        """Giữ bìa + khối mục lục của file mẫu, bỏ phần thân mẫu, Việt hoá bìa."""
        body = self.doc.element.body
        children = list(body.iterchildren())
        sdt = next((c for c in children if c.tag == qn('w:sdt')), None)
        if sdt is None:
            fail('file mẫu không có khối mục lục (w:sdt) sau bìa')
        for c in children[children.index(sdt) + 1:]:
            if c.tag != qn('w:sectPr'):
                body.remove(c)

        title = None
        for p in self.doc.paragraphs:
            t = p.text.strip()
            if t == 'Capstone Project Report':
                self.set_text(p, 'Tài liệu dự án')
            elif t.startswith('Report 6'):
                self.set_text(p, 'Report 6 – Hướng dẫn sử dụng phần mềm')
                title = p
            elif t.startswith('– Hanoi'):
                self.set_text(p, '– Hà Nội, Tháng %d năm %d –' % (today.month, today.year))
        if title is None:
            fail('không thấy dòng "Report 6 – ..." trên bìa file mẫu')
        name = Paragraph(title._p.getnext(), title._parent)   # đoạn trống ngay dưới
        r = name.add_run(PROJECT_NAME)
        r.bold, r.font.size = True, Pt(16)
        # Tên dự án chiếm hai dòng chỗ một dòng trống: bỏ bớt một dòng trống bên dưới
        # cho ngày tháng không bị đẩy sang trang 2.
        spare = name._p.getnext()
        if spare is not None and spare.tag == qn('w:p') and not Paragraph(spare, None).text.strip():
            spare.getparent().remove(spare)

        content = sdt.find(qn('w:sdtContent'))
        paras = content.findall(qn('w:p'))
        for extra in paras[1:]:
            content.remove(extra)
        head = Paragraph(paras[0], None)
        self.set_text(head, 'Mục lục')
        page_break_before(head)
        content.append(parse_xml(
            '<w:p %s><w:pPr><w:pStyle w:val="TOC1"/><w:tabs><w:tab w:val="right" w:leader="dot" w:pos="9040"/>'
            '</w:tabs></w:pPr><w:r><w:fldChar w:fldCharType="begin"/></w:r>'
            '<w:r><w:instrText xml:space="preserve"> TOC \\o "1-4" \\h \\z \\u </w:instrText></w:r>'
            '<w:r><w:fldChar w:fldCharType="separate"/></w:r>'
            '<w:r><w:t>Nhấn chuột phải vào đây rồi chọn “Update Field” để sinh mục lục.</w:t></w:r>'
            '<w:r><w:fldChar w:fldCharType="end"/></w:r></w:p>' % nsdecls('w')))

        # Số trang ở chân trang, trừ trang bìa.
        sec = self.doc.sections[0]
        sec.different_first_page_header_footer = True
        sec.first_page_footer.paragraphs[0].text = ''
        fp = sec.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fp._p.append(parse_xml('<w:fldSimple %s w:instr="PAGE"><w:r><w:t>1</w:t></w:r></w:fldSimple>'
                               % nsdecls('w')))

        # Kiểm chính tả theo tiếng Việt: file mẫu đặt en-GB/en-US, Word gạch đỏ gần hết các chữ.
        for lang in self.doc.styles.element.iter(qn('w:lang')):
            lang.set(qn('w:val'), 'vi-VN')
        props = self.doc.core_properties
        props.title = 'Report 6 – Hướng dẫn sử dụng phần mềm POSCS'
        props.subject = PROJECT_NAME
        props.author = 'Nhóm dự án POSCS'
        props.last_modified_by = 'tools/guide/word/export_word.py'
        props.revision = 1

    @staticmethod
    def set_text(p, text):
        runs = p.runs
        if not runs:
            p.add_run(text)
            return
        runs[0].text = text
        for r in runs[1:]:
            r._r.getparent().remove(r._r)

    def module(self, mod):
        self.heading('%s %s' % (mod.num, mod.title), 3, new_page=True)
        diagram = DIAGRAMS.get(mod.key)
        for node in mod.soup.children:
            if diagram and isinstance(node, Tag) and node.get('id') == diagram[2]:
                self.picture(self.doc, HERE / diagram[0], mod.diagram_fig, diagram[1], self.text_w)
            self.block(node, mod, self.doc, self.text_w)


# ----------------------------------------------------------------------
# Nội dung viết tay: lịch sử, bàn giao, cài đặt, tổng quan
# ----------------------------------------------------------------------

# (ngày, T/S/X, mô tả). Người phụ trách để trống: nhóm điền khi nộp.
CHANGES = [
    ('25/09/2026', 'T', 'Hướng dẫn sử dụng trong ứng dụng: Khách hàng, Hợp đồng (PR #151).'),
    ('25/09/2026', 'S', 'Hợp đồng: Xuất PDF chỉ có ở hợp đồng bán (PR #152).'),
    ('25/09/2026', 'T', 'Hướng dẫn Phiếu hỗ trợ (PR #153).'),
    ('26/09/2026', 'T', 'Hướng dẫn Nhân viên (PR #154).'),
    ('28/09/2026', 'S', 'Khách hàng: Sales đã được giao tỉnh chỉ sửa, xoá, đánh giá khách của mình (PR #155).'),
    ('28/09/2026', 'T', 'Hướng dẫn Sản phẩm (PR #156).'),
    ('28/09/2026', 'T', 'Hướng dẫn phần chung: đăng nhập, mật khẩu, trang chủ, thông báo, thông tin cá nhân '
                        '(PR #157).'),
    ('28/09/2026', 'T', 'Tài liệu này: dựng tự động từ hướng dẫn trong ứng dụng; thêm gói bàn giao, hướng dẫn '
                        'cài đặt, phần tổng quan.'),
    ('29/09/2026', 'S', 'Phiếu hỗ trợ, Sản phẩm, Khách hàng: xoá không thành công thì trang chi tiết mở lại kèm '
                        'thông báo (PR #159).'),
    ('29/09/2026', 'X', 'Hợp đồng: gỡ xuất PDF và nhập PDF theo yêu cầu khách hàng (PR #160).'),
    ('29/09/2026', 'S', 'Khách hàng: xoá từ danh sách Nhà cung cấp về đúng danh sách; danh sách báo khi không tìm thấy '
                        'bản ghi (PR #161).'),
]


def repo_facts():
    """Những con số lấy thẳng từ repo, để tài liệu không ghi lệch code."""
    migrations = [int(m.group(1)) for f in (ROOT / 'db' / 'migrations').glob('V*__*.sql')
                  for m in [re.match(r'V(\d+)__', f.name)] if m]
    schema = (ROOT / 'db' / 'schema.sql').read_text(encoding='utf-8')
    accounts = len(re.findall(r"^\('\w+', '\$2a\$", schema, flags=re.M))
    props = (ROOT / 'nbproject' / 'project.properties').read_text(encoding='utf-8')
    java = re.search(r'^javac\.target=(\S+)', props, flags=re.M).group(1)
    war = re.search(r'^war\.name=(\S+)', props, flags=re.M).group(1)
    jbcrypt = next((ROOT / 'lib').glob('jbcrypt-*.jar')).name
    # Hai jar NetBeans thường tự cấp khi build; ngoài IDE phải chỉ ra bằng -D (xem DEPLOY.md).
    servlet = next((ROOT / 'lib-build').glob('jakarta.servlet-api-*.jar')).name
    copylibs = next((ROOT / 'lib-build').glob('*copylibstask*.jar')).name
    if 'Admin' in re.findall(r"role_name = '(\w+)'\)", schema.split('INSERT IGNORE INTO users')[1][:20000]):
        fail('schema.sql đã gieo tài khoản Admin: sửa lại bước tạo Admin ở mục cài đặt')
    return {'v_max': max(migrations), 'accounts': accounts, 'java': java, 'war': war,
            'context': war[:-4], 'jbcrypt': jbcrypt, 'servlet': servlet, 'copylibs': copylibs}


def write_changes(w):
    w.heading('I. Lịch sử thay đổi', 1, new_page=True)
    w.plain_table(['Ngày', 'T*\nS, X', 'Người phụ trách', 'Mô tả thay đổi'],
                  [(d, k, '', desc) for d, k, desc in CHANGES], [2.2, 1.2, 2.6, 10])
    w.p('*T - Thêm, S - Sửa, X - Xoá')


def write_package(w, f):
    w.heading('II. Gói bàn giao và hướng dẫn sử dụng', 1, new_page=True)
    w.heading('1. Gói bàn giao', 2)
    w.p('Các hạng mục bàn giao của phiên bản này.')
    items = [
        ('Kế hoạch và theo dõi tiến độ dự án', ''),
        ('Product Backlog', ''),
        ('Mã nguồn', 'Kho GitHub %s, nhánh main. Ứng dụng web Java %s, Jakarta EE 11 (Servlet + JSP), dự án '
                     'Ant / NetBeans. Bản đóng gói: dist/%s (lệnh ant dist).' % (REPO_URL[:-4], f['java'], f['war'])),
        ('Script cơ sở dữ liệu', 'db/schema.sql: dựng cơ sở dữ liệu mới, kèm dữ liệu mẫu. db/migrations/V1 … V%d: '
                                 'các thay đổi tăng dần cho cơ sở dữ liệu đang chạy (cách dùng ở '
                                 'db/migrations/README.md).' % f['v_max']),
        ('Báo cáo cuối kỳ', ''),
        ('Tài liệu test case', 'Test tự động trong kho mã: test/ (JUnit 4 + Mockito; test tích hợp chạy trên MySQL '
                               'thật) và test/js/ (node --test). Tài liệu test case Excel dựng bằng tools/testdoc/.'),
        ('Danh sách lỗi (Defects)', ''),
        ('Danh sách vấn đề (Issues)', ''),
        ('Slide', ''),
        ('Hướng dẫn sử dụng', 'Trong ứng dụng: trang Hướng dẫn, 6 phần, mở từ nút Hướng dẫn trên các trang. '
                              'Tài liệu này dựng từ chính các trang đó bằng tools/guide/word/export_word.py.'),
    ]
    w.plain_table(['STT', 'Hạng mục bàn giao', 'Mô tả'],
                  [(str(i), a, b) for i, (a, b) in enumerate(items, 1)], [1.2, 4.4, 10.4])


def write_install(w, f, book):
    w.heading('2. Hướng dẫn cài đặt', 2)
    w.heading('2.1 Yêu cầu hệ thống', 3)
    w.plain_table(['Thành phần', 'Yêu cầu'], [
        ('Máy chủ ứng dụng', 'JDK %s trở lên. Apache Tomcat 11, hoặc máy chủ khác hỗ trợ Jakarta EE 11.' % f['java']),
        ('Cơ sở dữ liệu', 'MySQL 8.x trở lên (đã chạy trên MySQL 9.3), bảng mã utf8mb4.'),
        ('Công cụ đóng gói', 'Apache Ant, hoặc NetBeans. Mọi thư viện đi kèm sẵn: lib/ (chạy cùng ứng dụng) và '
                             'lib-build/ (chỉ để biên dịch).'),
        ('Thư mục lưu file', 'Một thư mục ghi được, nằm ngoài thư mục triển khai, chứa file người dùng tải lên '
                             '(ảnh, catalogue sản phẩm).'),
        ('Gửi email (không bắt buộc)', 'Tài khoản Gmail có App Password, để gửi mã OTP và thông tin tài khoản. '
                                       'Không cấu hình thì nội dung thư chỉ được ghi ra log máy chủ.'),
        ('Máy người dùng', 'Trình duyệt Chrome, Edge hoặc Firefox bản mới.'),
    ], [4.5, 11.5])

    w.heading('2.2 Các bước cài đặt', 3)
    w.step_title('Bước 1. Lấy mã nguồn')
    w.code(['git clone ' + REPO_URL])
    w.step_title('Bước 2. Tạo cơ sở dữ liệu')
    w.code(['mysql -u root -p -e "CREATE DATABASE poscs_db CHARACTER SET utf8mb4 COLLATE utf8mb4_unicode_ci;"',
            'mysql --default-character-set=utf8mb4 -u root -p poscs_db < db/schema.sql'])
    w.p('`db/schema.sql` dựng đủ mọi bảng, kèm dữ liệu mẫu: danh mục tỉnh / xã, sản phẩm, khách hàng, hợp đồng, '
        'phiếu hỗ trợ và %d tài khoản thử dùng chung mật khẩu `Poscs@123`. File này xoá rồi tạo lại mọi bảng, nên '
        'chỉ chạy trên cơ sở dữ liệu mới. Cơ sở dữ liệu đã có dữ liệu thì không chạy lại file này, mà áp các file '
        'mới trong `db/migrations/` theo thứ tự (xem `db/migrations/README.md`).' % f['accounts'])
    w.step_title('Bước 3. Đóng gói ứng dụng')
    w.p('Chạy ở thư mục mã nguồn:')
    # Ba dòng có dấu nối, chữ nhỏ: để liền một dòng thì Word ngắt ngay sau dấu "-" của "-D...".
    w.code(['ant -Dj2ee.platform.classpath=lib-build/%s \\' % f['servlet'],
            '    -Dlibs.CopyLibs.classpath=lib-build/%s \\' % f['copylibs'],
            '    dist'], size=8)
    w.p('Trên Windows (cmd), gõ liền thành một dòng và bỏ các dấu `\\` ở cuối dòng. Lệnh tạo file `dist/%s`. '
        'Hai tham số `-D` là thứ NetBeans vẫn tự cấp khi build: thư viện Servlet và tác vụ CopyLibs. Cả hai jar '
        'có sẵn trong `lib-build/`, chỉ dùng lúc biên dịch, không vào file WAR. Thiếu chúng, `ant dist` dừng với '
        'lỗi "The Java EE server classpath is not correctly set up". Dùng NetBeans thì mở thư mục dự án rồi chọn '
        'Clean and Build.' % f['war'])
    w.step_title('Bước 4. Khai báo cấu hình')
    w.p('Ứng dụng đọc cấu hình từ biến môi trường của tiến trình Tomcat. Đặt các biến trong `bin/setenv.sh` '
        '(Linux) hoặc `bin/setenv.bat` (Windows) của Tomcat:')
    w.plain_table(['Biến', 'Ý nghĩa', 'Bỏ trống thì'], [
        ('DB_URL', 'Chuỗi kết nối MySQL', 'jdbc:mysql://localhost:3306/poscs_db'),
        ('DB_USER', 'Tài khoản MySQL', 'root'),
        ('DB_PASSWORD', 'Mật khẩu MySQL', '1234'),
        ('UPLOAD_DIR', 'Thư mục lưu file tải lên', '<thư mục người dùng>/poscs_uploads'),
        ('MAIL_USERNAME', 'Địa chỉ Gmail gửi thư', 'Không gửi thư, nội dung thư ghi ra log'),
        ('MAIL_PASSWORD', 'App Password 16 ký tự của Gmail (không phải mật khẩu đăng nhập Gmail)', 'Như trên'),
        ('MAIL_SMTP_HOST, MAIL_SMTP_PORT, MAIL_FROM', 'Chỉ đặt khi không gửi qua Gmail',
         'smtp.gmail.com, 587, bằng MAIL_USERNAME'),
    ], [4.2, 5.6, 6.2])
    w.p('Giá trị mặc định chỉ dành cho máy phát triển. Ví dụ `bin/setenv.sh`:')
    w.code(['export DB_URL="jdbc:mysql://<máy chủ CSDL>:3306/poscs_db"',
            'export DB_USER="poscs_app"',
            'export DB_PASSWORD="<mật khẩu mạnh>"',
            'export UPLOAD_DIR="/var/lib/poscs/uploads"',
            'export MAIL_USERNAME="<địa chỉ Gmail>"',
            'export MAIL_PASSWORD="<App Password>"'])
    w.p('Trên Windows, mỗi dòng của `bin/setenv.bat` viết dạng `set "DB_USER=poscs_app"`.')
    w.step_title('Bước 5. Triển khai và khởi động')
    w.p('Chép `dist/%s` vào thư mục `webapps/` của Tomcat rồi khởi động Tomcat (`bin/startup.sh` hoặc '
        '`bin/startup.bat`). Mở `http://<máy chủ>:8080/%s/`: trang đăng nhập hiện ra.' % (f['war'], f['context']))
    w.step_title('Bước 6. Tạo tài khoản Admin đầu tiên')
    w.p('Dữ liệu mẫu cố ý không có tài khoản Admin: mã nguồn công khai, nên mật khẩu gieo sẵn nào cũng coi như ai '
        'cũng biết. Trang Nhân viên cũng không cho chọn vai trò Admin. Admin đầu tiên tạo thẳng trong cơ sở dữ '
        'liệu. Sinh chuỗi băm của một mật khẩu tạm bằng thư viện đi kèm (`jshell` có sẵn trong JDK), chạy ở thư '
        'mục mã nguồn:')
    w.code(['jshell --class-path lib/%s' % f['jbcrypt'],
            'jshell> org.mindrot.jbcrypt.BCrypt.hashpw("<mật khẩu tạm>", org.mindrot.jbcrypt.BCrypt.gensalt())'])
    w.p('Chép chuỗi bắt đầu bằng `$2a$10$` vừa in ra, rồi chạy trong MySQL (thay họ, tên):')
    w.code(['INSERT INTO users (username, password_hash, must_change_password, role_id,',
            '                   last_name, first_name, department_id, hire_date)',
            "VALUES ('admin', '<chuỗi $2a$10$… vừa chép>', 1,",
            "        (SELECT role_id FROM roles WHERE role_name = 'Admin'),",
            "        '<Họ>', '<Tên>',",
            "        (SELECT department_id FROM departments WHERE department_name = 'Ban giám đốc'),",
            '        CURDATE());'])
    w.p('`must_change_password = 1` buộc đổi mật khẩu tạm này ngay lần đăng nhập đầu.')
    w.step_title('Bước 7. Đăng nhập lần đầu')
    w.p('Đăng nhập bằng tài khoản vừa tạo. Hệ thống buộc đổi mật khẩu tạm, rồi bổ sung thông tin cá nhân (mục %s). '
        'Sau đó tạo tài khoản cho nhân viên ở trang Nhân viên (mục %s).'
        % (book.section('general', 'lan-dau'), book.section('employee', 'them-nhan-vien')))
    w.step_title('Trước khi chạy thật')
    w.bullet_list([
        'Tạo tài khoản MySQL riêng cho ứng dụng, chỉ có quyền SELECT, INSERT, UPDATE, DELETE trên `poscs_db`; '
        'không dùng `root`:',
    ])
    w.code(["CREATE USER 'poscs_app'@'%' IDENTIFIED BY '<mật khẩu mạnh>';",
            "GRANT SELECT, INSERT, UPDATE, DELETE ON poscs_db.* TO 'poscs_app'@'%';",
            'FLUSH PRIVILEGES;'])
    w.bullet_list([
        'Không mở cổng MySQL (3306) ra Internet; chỉ máy chủ ứng dụng kết nối được.',
        'Bật HTTPS: qua reverse proxy (ví dụ Nginx) hoặc connector SSL của Tomcat.',
        'Khoá các tài khoản thử của dữ liệu mẫu ở trang Nhân viên (mục %s): mật khẩu chung của chúng nằm công '
        'khai trong mã nguồn. Dữ liệu mẫu khác (khách hàng, hợp đồng, phiếu) cũng đi kèm `db/schema.sql`; kho mã '
        'chưa có script dựng cơ sở dữ liệu trống.' % book.section('employee', 'khoa'),
        'Sao lưu cơ sở dữ liệu định kỳ bằng `mysqldump`.',
    ])


def write_overview(w, book):
    w.heading('3. Hướng dẫn sử dụng', 2, new_page=True)
    w.heading('3.1 Tổng quan', 3)
    w.p('POSCS là ứng dụng web nội bộ của POSTEF – Chi nhánh Miền Bắc, dùng để quản lý khách hàng (khách mua và nhà '
        'cung cấp), hợp đồng bán và mua, danh mục sản phẩm và phiếu hỗ trợ kỹ thuật. Trang chủ tổng hợp số liệu; '
        'hệ thống tự gửi thông báo khi hợp đồng sắp hết hạn hoặc phiếu hỗ trợ sắp quá hạn xử lý.')
    w.picture(w.doc, HERE / OVERVIEW[0], 1, OVERVIEW[1], w.text_w)
    w.p('Mỗi tài khoản có một trong ba vai trò. Vai trò quyết định trang nào mở được và nút nào hiện ra:')
    w.plain_table(['Vai trò', 'Làm được', 'Chỉ xem, hoặc không vào được'], [
        ('Admin', 'Mọi việc, gồm quản lý nhân viên: tạo tài khoản, giao tỉnh phụ trách, khoá tài khoản.', '—'),
        ('Sales', 'Khách hàng, hợp đồng, phiếu hỗ trợ.',
         'Chỉ xem sản phẩm; không vào trang Nhân viên. Sales đã có cấp trên trong sơ đồ tổ chức chỉ xem khách hàng '
         'và hợp đồng.'),
        ('Kỹ thuật', 'Danh mục sản phẩm. Phần xử lý của phiếu hỗ trợ được giao cho mình: trạng thái, nguyên nhân, '
                     'phương hướng xử lý, kết quả.',
         'Chỉ xem khách hàng, hợp đồng và phần còn lại của phiếu; không vào trang Nhân viên.'),
    ], [2.6, 6.7, 6.7])
    w.p('Cách đọc phần hướng dẫn:')
    w.bullet_list([
        'Mục %s đến %s đi đúng thứ tự các tab của trang Hướng dẫn trong ứng dụng: %s.'
        % (book.modules[0].num, book.modules[-1].num, ', '.join(m.tab for m in book.modules)),
        'Dòng **Ai làm được** cho biết vai trò nào làm được việc đó; vai trò bị gạch ngang là không làm được.',
        'Số tròn đỏ ở các bước ứng với số khoanh đỏ trên hình ngay phía trên. Số tròn xám là bước không khoanh '
        'trên hình.',
        'Trong ứng dụng, nút **Hướng dẫn** trên các trang mở thẳng mục tương ứng; trang Hướng dẫn có mục lục và '
        'nút In / Lưu PDF.',
    ])


def main():
    ap = argparse.ArgumentParser(description='Dựng Report 6 (Word) từ hướng dẫn trong app.')
    ap.add_argument('--template', required=True, help='file mẫu Report6_Software User Guides.docx')
    ap.add_argument('--out', required=True, help='file .docx sẽ ghi ra')
    args = ap.parse_args()
    template, out = Path(args.template), Path(args.out)
    if out.resolve() == template.resolve():
        fail('--out trùng --template: file mẫu phải giữ nguyên')

    book = Book()
    facts = repo_facts()
    w = Writer(template, book)
    w.prepare(datetime.date.today())
    write_changes(w)
    write_package(w, facts)
    write_install(w, facts, book)
    write_overview(w, book)
    for mod in book.modules:
        w.module(mod)
    w.doc.save(str(out))
    # In không dấu: console Windows (cp1252) không in được tiếng Việt.
    print('xong: %d phan, %d hinh -> %s' % (len(book.modules), book.fig_total, out))


if __name__ == '__main__':
    main()
